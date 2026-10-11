"""Read synthetic fixtures and package first-use studies. No model calls or XLSX edits."""
from collections import Counter, defaultdict
from decimal import Decimal
from pathlib import Path
import csv
import hashlib
import json
import sqlite3
import zipfile

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
GUIDES = ROOT / "evaluation" / "user_test_round1"
OUTPUT = ROOT / "runtime" / "user-test-round1-20261011"
EXPECTED = [(16, 14, 4035), (16, 13, 6370), (15, 14, 11751), (18, 16, 8570), (17, 15, 9510)]
GUIDE_NAMES = ["participant_count.md", "participant_compare.md", "participant_duplicate.md", "participant_compare.md", "participant_join.md"]


def workbook_reference(path):
    with path.open("rb") as source:
        wb = openpyxl.load_workbook(source, read_only=True, data_only=True)
        sheet = wb.worksheets[0]
        values = list(sheet.values)
        headers = list(values[0])
        rows = [dict(zip(headers, row)) for row in values[1:] if any(value is not None for value in row)]
        ids = Counter(str(row["order_id"]).strip() for row in rows if row.get("order_id"))
        amount = Decimal(0)
        groups = defaultdict(Decimal)
        excluded_group_amount = Decimal(0)
        for row in rows:
            value = row.get("amount")
            if not isinstance(value, (int, float)):
                continue
            numeric = Decimal(str(value))
            amount += numeric
            if row.get("region"):
                groups[str(row["region"])] += numeric
            else:
                excluded_group_amount += numeric
        result = {
            "file": path.name, "sheet": sheet.title, "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
            "rows": len(rows), "distinct_nonempty_order_ids": len(ids),
            "duplicate_order_ids": {key: value for key, value in ids.items() if value > 1},
            "exact_duplicate_groups": sum(value > 1 for value in Counter(tuple(row.get(key) for key in headers) for row in rows).values()),
            "all_numeric_amount_sum": str(amount),
            "raw_group_amount_sums": {key: str(value) for key, value in groups.items()},
            "amount_with_missing_region": str(excluded_group_amount),
            "missing_region_rows": sum(not row.get("region") for row in rows),
            "missing_amount_rows": sum(row.get("amount") is None for row in rows),
            "scope": "Original rows, all statuses, no cleanup. Amount sums are arithmetic controls, not revenue definitions.",
        }
        assert sum(groups.values()) + excluded_group_amount == amount
        wb.close()
    return result


def join_reference():
    db = sqlite3.connect(":memory:")
    db.execute("CREATE TABLE orders (order_id TEXT, customer_id TEXT, order_date TEXT, amount INTEGER)")
    db.execute("CREATE TABLE items (order_id TEXT, item_id TEXT, category TEXT)")
    for table, name in [("orders", "join_orders_sample.csv"), ("items", "join_items_sample.csv")]:
        with (ROOT / "evaluation" / "fixtures" / name).open(encoding="utf-8-sig", newline="") as source:
            reader = csv.reader(source)
            next(reader)
            placeholders = ",".join("?" for _ in range(4 if table == "orders" else 3))
            db.executemany(f"INSERT INTO {table} VALUES ({placeholders})", list(reader))
    source = db.execute("SELECT COUNT(*),COUNT(DISTINCT order_id),SUM(amount) FROM orders").fetchone()
    left = db.execute("SELECT COUNT(*),COUNT(DISTINCT o.order_id),SUM(o.amount) FROM orders o LEFT JOIN items i ON o.order_id=i.order_id").fetchone()
    inner = db.execute("SELECT COUNT(*),COUNT(DISTINCT o.order_id),SUM(o.amount) FROM orders o JOIN items i ON o.order_id=i.order_id").fetchone()
    db.close()
    assert source == (4, 4, 410) and left == (5, 4, 530) and inner == (4, 3, 470)
    return {"source_rows_distinct_orders_amount": source, "left_join": left, "inner_join": inner,
            "scope": "Independent SQLite controls; duplicated source amount must block unsafe sum even if row counts or net changes seem plausible."}


def main():
    if OUTPUT.exists():
        raise SystemExit(f"Output already exists; preserve it and choose a new dated directory: {OUTPUT}")
    references = []
    for index, expected in enumerate(EXPECTED, 1):
        path = ROOT / "evaluation" / "novice_test_files" / f"novice_test_p{index:02}.xlsx"
        result = workbook_reference(path)
        assert (result["rows"], result["distinct_nonempty_order_ids"], Decimal(result["all_numeric_amount_sum"])) == expected
        references.append(result)
    joins = join_reference()
    OUTPUT.mkdir(parents=True)
    archives = []
    for index, guide_name in enumerate(GUIDE_NAMES, 1):
        excel = ROOT / "evaluation" / "novice_test_files" / f"novice_test_p{index:02}.xlsx"
        archive = OUTPUT / f"MetricGround-P{index:02}-participant.zip"
        with zipfile.ZipFile(archive, "x", zipfile.ZIP_DEFLATED) as bundle:
            bundle.write(GUIDES / guide_name, "任务说明.md")
            bundle.write(excel, excel.name)
            if index == 5:
                for name in ["join_orders_sample.csv", "join_items_sample.csv"]:
                    bundle.write(ROOT / "evaluation" / "fixtures" / name, name)
        with zipfile.ZipFile(archive) as bundle:
            assert bundle.testzip() is None
            assert hashlib.sha256(bundle.read(excel.name)).hexdigest() == references[index - 1]["sha256"]
            assert set(bundle.namelist()) == ({"任务说明.md", excel.name} | ({"join_orders_sample.csv", "join_items_sample.csv"} if index == 5 else set()))
        archives.append(str(archive))
    (OUTPUT / "host_reference_private.json").write_text(json.dumps({
        "software_baseline": "v0.3.9 / Sites13 / eb90c7925ccbd04148479a6534298d5eec93c371",
        "participants_completed": None, "test_status": "not_started", "fixtures": references, "join_controls": joins,
    }, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"packages": archives, "reference_checks_passed": True, "participant_results": "not_started", "model_requests": 0}, ensure_ascii=True, indent=2))


if __name__ == "__main__":
    main()
